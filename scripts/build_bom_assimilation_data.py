"""Read-only assimilation of BOM sources and the manually reviewed mapping sheets."""
from __future__ import annotations
import json, re, sys
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"D:\CSD Software\Store\BOMs")
OUT = Path("docs/bom-assimilation-data.json")
DECISIONS = Path("bom-assimilation/review-decisions.json")
sys.path.insert(0, r"D:\CSD Software\Store\api-client")
from store_client import ClientConfig, StoreClient  # noqa: E402

def clean(value): return str(value or "").strip()
def norm(value): return re.sub(r"[^a-z0-9]+", " ", clean(value).lower()).strip()
def code_tokens(value): return re.findall(r"(?:IMP|CMP)-[A-Z0-9-]+", clean(value).upper())
def provenance(value):
    records=[]
    for line in clean(value).splitlines():
        m=re.match(r"(.+) · (.+) · row (\d+)$", line)
        if m: records.append((m.group(1),m.group(2),int(m.group(3))))
    return records

RAW_CACHE = {}
def raw_record(file_rel, sheet_name, row_number):
    cache_key=(file_rel,sheet_name)
    if cache_key in RAW_CACHE:
        header_row, all_rows = RAW_CACHE[cache_key]
        return {clean(header): clean(value) for header,value in zip(header_row,all_rows[row_number-1]) if clean(header) and clean(value)}
    file = ROOT / file_rel
    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb[sheet_name]
    header_row=None
    for row in ws.iter_rows(values_only=True):
        if any(re.search(r"comment|item name|part no|part name|connector|^item$", str(cell or ""), re.I) for cell in row):
            header_row=row; break
    if not header_row: return {}
    all_rows=list(ws.iter_rows(values_only=True))
    RAW_CACHE[cache_key]=(header_row,all_rows)
    values=all_rows[row_number-1]
    return {clean(header): clean(value) for header,value in zip(header_row,values) if clean(header) and clean(value)}

def load_review(path, exact_mpn=False):
    ws=load_workbook(path,data_only=True)["Components"]
    headers=[clean(ws.cell(4,c).value) for c in range(1,ws.max_column+1)]
    pos={name:i+1 for i,name in enumerate(headers)}; rows=[]
    for row in range(5,ws.max_row+1):
        title=clean(ws.cell(row,pos['Component']).value)
        if not title: continue
        manual=clean(ws.cell(row,pos['Store Item Code']).value)
        recommendation=clean(ws.cell(row,pos['Recommended Code']).value)
        effective = recommendation if exact_mpn else manual
        rows.append({
          'title':title,'specification':clean(ws.cell(row,pos['Specification']).value),
          'mpn':clean(ws.cell(row,pos['Manufacturer Part Number']).value), 'supplier':clean(ws.cell(row,pos['Supplier / LCSC Number']).value),
          'manualDecision':manual, 'recommendedCode':recommendation, 'effectiveDecision':effective,
          'matchStrength':clean(ws.cell(row,pos['Match Strength']).value), 'matchReason':clean(ws.cell(row,pos['Recommendation Reason']).value),
          'notes':clean(ws.cell(row,pos.get('Column1',0)).value) if pos.get('Column1') else '',
          'provenance':provenance(ws.cell(row,pos['Source Provenance']).value),
          'sourceRows':clean(ws.cell(row,pos['Source Rows']).value), 'sourceFiles':clean(ws.cell(row,pos['Source Files']).value),
        })
    return rows

def context_name(row):
    title, spec = row['title'], row['specification']
    generic={'pcb','main pcb','connector','metal connector','wire','heat shrink','acrylic','3d print','components'}
    if norm(title) in generic or norm(title).endswith(' pcb'):
        first=row['provenance'][0][0] if row['provenance'] else ''
        parts=[part for part in re.split(r"[\\/]", first) if part]
        project=' '.join(parts[-4:-1]) if len(parts)>=4 else ''
        label=' '.join(part for part in [project,title,spec] if part)
        return re.sub(r"\s+", " ", label).strip()
    return ' — '.join(part for part in [title,spec] if part)

def source_key(file_rel, sheet_name, row_number):
    return (file_rel, sheet_name, int(row_number))

def load_decisions():
    payload=json.loads(DECISIONS.read_text(encoding='utf-8'))
    mapping={}; expansions={}
    for decision in payload['decisions']:
        if decision['kind']=='SOURCE_EXPANSION':
            source=decision['source']; expansions[source_key(source['file'],source['sheet'],source['row'])]=(decision,decision['outputs'])
        for entry in decision.get('mappings',[]):
            source=entry['source']; mapping[source_key(source['file'],source['sheet'],source['row'])]=(decision,entry['output'])
    return payload, mapping, expansions

def default_output(row, by_code, exception_index):
    decision=row['effectiveDecision']; lower=norm(decision); codes=code_tokens(decision)
    if row['manualDecision'] and ('ignore' in lower or 'mistake' in lower):
        return {'canonicalKey':'EXCLUDED:'+norm(row['title'])+'|'+norm(row['specification']), 'disposition':'EXCLUDED'}
    if len(codes)==1 and codes[0] in by_code and not lower.startswith('new'):
        return {'canonicalKey':'STORE:'+codes[0], 'disposition':'EXISTING_STORE_ITEM', 'storeItemCode':codes[0]}
    if len(codes)>1 or '/' in decision or 'same as above' in lower:
        return {'canonicalKey':'REVIEW:'+norm(row['title'])+'|'+norm(row['specification'])+'|'+str(exception_index), 'disposition':'REVIEW_DIRECTIVE'}
    key='NEW:'+norm(row['mpn']) if row['mpn'] else 'NEW:'+norm(row['supplier']) if row['supplier'] else 'NEW:'+norm(row['title'])+'|'+norm(row['specification'])
    return {'canonicalKey':key, 'disposition':'NEW_COMPONENT_CANDIDATE'}

def main():
    client=StoreClient(ClientConfig.load(Path(r"D:\CSD Software\Store\api-client\.env.local")))
    items=[]; page=1
    while True:
        result=client.request('GET','/api/v1/items',query=[('page',str(page)),('pageSize','200')]); items.extend(result['items'])
        if page>=result['pagination']['totalPages']: break
        page+=1
    by_code={item['code'].upper():item for item in items}
    rows=load_review(Path('docs/bom-exact-mpn-recommendations.xlsx'),True)+load_review(Path('docs/bom-reconciliation-audit-remaining.xlsx'))
    decisions, reviewed_mapping, reviewed_expansions=load_decisions()
    groups=defaultdict(list); exceptions=[]; source_records=[]
    for row in rows:
        for file,sheet,row_number in row['provenance']:
            key_tuple=source_key(file,sheet,row_number)
            if key_tuple in reviewed_expansions:
                reviewed_decision, outputs=reviewed_expansions[key_tuple]
            elif key_tuple in reviewed_mapping:
                reviewed_decision, output=reviewed_mapping[key_tuple]; outputs=[output]
            else:
                reviewed_decision=None; outputs=[default_output(row,by_code,len(exceptions))]
            for output in outputs:
                key=output['canonicalKey']; disposition=output['disposition']
                if disposition=='REVIEW_DIRECTIVE':
                    exceptions.append({'component':row['title'],'decision':row['effectiveDecision'],'notes':row['notes'],'reason':'Multiple, relative, or mixed mapping directive requires explicit resolution.'})
                member={**row,'key':key,'disposition':disposition,'provenance':[(file,sheet,row_number)],'reviewDecisionId':reviewed_decision['id'] if reviewed_decision else None,'suggestedNameOverride':output.get('suggestedName',''),'quantityPerSetOverride':output.get('quantityPerSet','')}
                groups[key].append(member)
                raw=raw_record(file,sheet,row_number)
                source_records.append({'canonicalKey':key,'disposition':disposition,'component':row['title'],'file':file,'sheet':sheet,'row':row_number,'rawFields':raw,'reviewDecisionId':member['reviewDecisionId'],'quantityPerSetOverride':member['quantityPerSetOverride']})
    canonical=[]
    for key,members in groups.items():
        first=members[0]; code=key[6:] if key.startswith('STORE:') else ''
        item=by_code.get(code)
        canonical.append({'canonicalKey':key,'disposition':first['disposition'],'suggestedName':first['suggestedNameOverride'] or (item['title'] if item else context_name(first)),'storeItemCode':code,'storeItemTitle':item['title'] if item else '', 'mpns':sorted({r['mpn'] for r in members if r['mpn']}),'supplierParts':sorted({r['supplier'] for r in members if r['supplier']}),'sourceIdentityCount':len(members),'sourceOccurrenceCount':len(members),'manualDecisions':sorted({r['manualDecision'] for r in members if r['manualDecision']}),'reviewDecisionIds':sorted({r['reviewDecisionId'] for r in members if r['reviewDecisionId']}),'notes':sorted({r['notes'] for r in members if r['notes']}),'sourceRecords':[record for record in source_records if record['canonicalKey']==key]})
    OUT.write_text(json.dumps({'catalogueItems':len(items),'canonical':canonical,'sourceRecords':source_records,'exceptions':exceptions,'reviewedDecisions':decisions['decisions']},indent=2),encoding='utf-8')
    print(json.dumps({'catalogueItems':len(items),'canonical':len(canonical),'existing':sum(x['disposition']=='EXISTING_STORE_ITEM' for x in canonical),'new':sum(x['disposition']=='NEW_COMPONENT_CANDIDATE' for x in canonical),'review':sum(x['disposition']=='REVIEW_DIRECTIVE' for x in canonical),'excluded':sum(x['disposition']=='EXCLUDED' for x in canonical),'sourceRecords':len(source_records),'exceptions':len(exceptions)}))
if __name__=='__main__': main()
